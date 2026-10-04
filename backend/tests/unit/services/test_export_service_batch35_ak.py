# TARGET-MODULE: backend.services.export_service
"""Battery AK - campaign #35 kill battery for backend/services/export_service.py.

Covers the 215 keys the CURRENT shipped suite (test_export_service.py, 121 tests)
leaves GREEN - a shipped-only per-key sweep proved the pool FRESH
(/home/agent/runs/c35-shipped-sweep.txt: "SHIPPED-SWEEP DONE 0 red / 215 green /
0 other"), so no stale-era keys are re-authored and none are skipped.

Killing machinery, by family:
  * Path.write_text SPY pinning kwargs EXACTLY ({"encoding": "utf-8"}) alongside
    byte-equal content: the sandbox forces UTF-8 mode (sys.flags.utf8_mode==1)
    so the BYTES of "utf-8"/"UTF-8"/None/dropped are identical, but the call
    itself is observable - all 19 encoding mutants RED on the kwargs pin.
  * a fake zipfile module swapped over es.zipfile: pins the ZipFile ctor
    ("w", ZIP_DEFLATED positionally, no kwargs) and every writestr byte.
  * a statement-signature session fake: every execute() is pinned by compiled
    Postgres text (literal_binds); UNPINNED sql raises, so select(None),
    where(None), == <-> != and execute(None) all RED.
  * exact update_progress / report_progress CALL-SEQUENCE spies ((args, kwargs)
    tuples, full expected list) - the 100-row tick can only land at idx 99/199.
  * a real-logging Handler on es.logger: exact record msg PLUS the exact
    extra= key set (rename/UPPER/drop/None all change it).
  * the TZ lever (TZ=Etc/GMT+11 + time.tzset(), restored in __exit__): filename
    timestamps must sit within 2 s of datetime.now(UTC); now(None) misses by
    11 h at every possible hour.
  * openpyxl READBACK pins: reload the saved bytes and assert header Font/Fill/
    Alignment/Border (4 sides, exact rgb strings - openpyxl stores color
    literals VERBATIM so 4472c4 != 4472C4), per-column widths, alt-row fills on
    rows 2/4 but not 3, freeze_panes == "A2"; raw sheet XML for "" vs None
    (value readback is blind: "" reloads as None, XML shows t="inlineStr").
  * width-family discriminating columns: a seed-0 column (display name "") with
    a falsy value pins width == 2.0 (the else-1 mutant stores 3.0), and a
    55-char header pinning a 55-char body TIE above the cap pins width 57.0
    (the >= mutant clamps the tie to min(55,50)=50 -> 52.0).
  * Z-date family by MESSAGE + polarity: dt.fromisoformat parses a trailing Z
    natively on 3.14, so the discriminator is a MID-STRING Z ("2026Z-01-05")
    whose ValueError message carries the REPLACED text "2026+00:00-01-05"
    (shipped) vs untouched "2026Z-01-05" (XXZXX); the lowercase-z twins are
    killed because shipped RAISES on "...T00:00:00z" while replace("z") parses.
  * del-attribute discriminator for the 2-arg getattr family: a dataclass
    field deleted after construction reads None under shipped 3-arg getattr
    and RAISES AttributeError under the default-dropped mutant.
  * mkdir polarity on a fresh nested EXPORT_DIR patch: parents True -> False /
    None makes FileNotFoundError, exist_ok drop on a pre-existing dir makes
    FileExistsError.

HONESTY LEDGER (registered equivalents - each PROVEN by construction probe,
see /home/agent/runs/c35-ledger.md + c35-probe/probe5-ledger.py, NOT inferred
from diff shape):
  * events_to_excel__mutmut_144  `len(str(value)) if value else 0` ->
        `if (value) or True else 0`: value is format_export_value's return,
        ALWAYS a str, and len(str("")) == 0 == the else-branch value, so the
        ternary is identical on every possible input.
Earlier EQUIV drafts DEMOTED to killable after probing (recorded so the
attribution trail is honest): the 19 encoding keys (write_text kwargs spy),
events_to_excel m145 (seed-0 column kills: 3.0 != 2.0), m146 (above-cap tie
kills: 52.0 != 57.0), m124 "" -> None (raw XML t= attribute), the 4 2-arg
getattr drops (del-field probe), get_selected_columns m6 `or True` (all-
invalid selection returns [] instead of EXTENDED), the 6 "UTF-8" alias keys
(the codec IS identical - but the CALL literal is observable), and the 4
XXZXX date keys (message-carrying ValueError).
"""

import asyncio
import csv
import io
import json
import logging
import os
import re
import tempfile
import time
import zipfile
from datetime import UTC, datetime
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import openpyxl
from sqlalchemy.dialects import postgresql

import backend.services.export_service as es
from backend.services.export_service import (
    EXTENDED_EXPORT_COLUMNS,
    DetectionExportRow,
    EventExportRow,
    ExportFormat,
    ExportService,
    detections_to_json,
    events_to_excel,
    filter_row_to_dict,
    format_export_value,
    generate_export_filename,
    get_selected_columns,
    parse_accept_header,
)

MOD = "backend.services.export_service"
T0 = datetime(2026, 1, 5, 10, 30, 45, tzinfo=UTC)
T1 = datetime(2026, 1, 5, 11, 30, 45, tzinfo=UTC)
TS_RE = r"\d{8}_\d{6}"


def run(coro):
    return asyncio.run(coro)


def assert_raises(exc_type, fn):
    try:
        fn()
    except exc_type as exc:
        return exc
    except Exception as other:
        raise AssertionError(f"wrong exception {type(other).__name__}: {other}") from other
    raise AssertionError(f"{exc_type.__name__} not raised")


def norm_sql(stmt, literal_binds=False):
    kw = {"literal_binds": True} if literal_binds else {}
    compiled = str(stmt.compile(dialect=postgresql.dialect(), compile_kwargs=kw))
    return " ".join(compiled.split())


# ---------------------------------------------------------------------------
# spies
# ---------------------------------------------------------------------------
class SpyFile:
    def __init__(self, name):
        self.name = name
        self.texts = []
        self.buffers = []

    def write_text(self, content, **kwargs):
        self.texts.append((content, kwargs))
        return len(content)

    def write_bytes(self, data):
        self.buffers.append(data)
        return len(data)

    def stat(self):
        if self.texts:
            return SimpleNamespace(st_size=len(self.texts[-1][0].encode("utf-8")))
        if self.buffers:
            return SimpleNamespace(st_size=len(self.buffers[-1]))
        return SimpleNamespace(st_size=0)

    def last_text(self):
        assert len(self.texts) == 1, f"expected one write_text, got {self.texts}"
        return self.texts[0]


class SpyDir:
    def __init__(self):
        self.files = {}

    def __truediv__(self, name):
        return self.files.setdefault(str(name), SpyFile(str(name)))

    def mkdir(self, **kwargs):
        pass

    def sole(self):
        assert len(self.files) == 1, f"expected one file, got {sorted(self.files)}"
        return next(iter(self.files.values()))


class ZipRec:
    def __init__(self):
        self.ctor = None
        self.entries = []


class FakeZipHandle:
    def __init__(self, rec, *args, **kwargs):
        self.rec = rec
        rec.ctor = (args, kwargs)

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False

    def writestr(self, name, data):
        self.rec.entries.append((name, data))


class FakeZipModule:
    ZIP_DEFLATED = 8

    def __init__(self):
        self.records = []

    def ZipFile(self, _fp, *args, **kwargs):
        rec = ZipRec()
        self.records.append(rec)
        return FakeZipHandle(rec, *args, **kwargs)


class FakeResult:
    def __init__(self, scalar=None, rows=None):
        self._scalar = scalar
        self._rows = rows or []

    def scalar(self):
        return self._scalar

    def scalars(self):
        return self

    def all(self):
        return list(self._rows)


class FakeSession:
    """Pins every execute() by compiled Postgres text; UNPINNED sql is fatal."""

    def __init__(self, pins):
        self.pins = pins
        self.executed = []

    async def execute(self, stmt):
        sql = norm_sql(stmt, literal_binds=True)
        self.executed.append(sql)
        for subs, scalar, rows in self.pins:
            if all(s in sql for s in subs):
                return FakeResult(scalar=scalar, rows=rows)
        raise AssertionError(f"UNPINNED SQL: {sql[:260]}")


class Tracker:
    def __init__(self):
        self.calls = []


def _update_progress(self, *args, **kwargs):
    self.calls.append((args, kwargs))


Tracker.update_progress = _update_progress


class Reporter:
    job_id = "job-ws-1"

    def __init__(self, start_raises=None):
        self.events = []
        self.fails = []
        self._start_raises = start_raises

    @property
    def duration_seconds(self):
        return 1.25

    async def start(self, metadata=None):
        if self._start_raises is not None:
            raise self._start_raises
        self.events.append(("start", metadata))

    async def report_progress(self, pct, current_step=None, force=False):
        self.events.append(("progress", pct, current_step, force))

    async def complete(self, result_summary=None):
        self.events.append(("complete", result_summary))

    async def fail(self, exc, retryable=None):
        self.fails.append((exc, retryable))


class LogCap(logging.Handler):
    def __init__(self):
        super().__init__()
        self.records = []

    def emit(self, record):
        self.records.append(record)


_STANDARD_LOG_KEYS = set(vars(logging.LogRecord("x", 0, "f", 1, "m", None, None))) | {
    "message",
    "asctime",
    "taskName",
}


def log_extras(record):
    return {k: vars(record)[k] for k in vars(record) if k not in _STANDARD_LOG_KEYS}


class TzPinned:
    """Put local time 11 h from UTC so datetime.now(None) != now(UTC) always."""

    def __enter__(self):
        self._saved = os.environ.get("TZ")
        os.environ["TZ"] = "Etc/GMT+11"
        time.tzset()
        return self

    def __exit__(self, *exc):
        if self._saved is None:
            os.environ.pop("TZ", None)
        else:
            os.environ["TZ"] = self._saved
        time.tzset()
        return False


def check_utc_timestamp(stamp):
    parsed = datetime.strptime(stamp, "%Y%m%d_%H%M%S").replace(tzinfo=UTC)
    drift = abs((datetime.now(UTC) - parsed).total_seconds())
    assert drift <= 2.0, f"timestamp {stamp} sits {drift}s from UTC now"


def filename_stamp(name, ext):
    m = re.fullmatch(rf"events_export_({TS_RE})\.{ext}", name)
    assert m, f"filename shape wrong: {name!r}"
    return m.group(1)


def ev(**kw):
    base = {
        "id": 1,
        "camera_id": "cam-a",
        "started_at": T0,
        "ended_at": T1,
        "risk_score": 87,
        "risk_level": "high",
        "summary": "=SUM(1)",
        "detection_count": 3,
        "reviewed": True,
        "object_types": "person,car",
        "reasoning": "because",
    }
    base.update(kw)
    return SimpleNamespace(**base)


BASE_ROW = [
    "1",
    "Front Door",
    T0.isoformat(),
    T1.isoformat(),
    "87",
    "high",
    "'=SUM(1)",
    "3",
    "Yes",
    "person,car",
    "because",
]


def unknown_row(i):
    return [
        str(i),
        "Unknown",
        T0.isoformat(),
        T1.isoformat(),
        "87",
        "high",
        "'=SUM(1)",
        "3",
        "Yes",
        "person,car",
        "because",
    ]


def three_event_pins():
    rows = [ev(id=1), ev(id=2, camera_id=None), ev(id=3, camera_id="cam-b")]
    pins = [
        (("count(*)",), 3, []),
        (("SELECT cameras.name", "cameras.id = 'cam-a'"), "Front Door", []),
        (("SELECT cameras.name", "cameras.id = 'cam-b'"), None, []),
        (("FROM events", "ORDER BY events.started_at DESC"), None, rows),
    ]
    return pins


class LoggerSwap:
    """Swap in a PLAIN logging.Logger: the app logger injects context keys
    (app_version/container_id/...) into every record, which would blur the
    exact extra= key-set assertion. Plain logger -> extras are exactly the
    keys the call passed."""

    def __init__(self):
        self.cap = LogCap()
        self.plain = logging.Logger("c35ak-plain")
        self.plain.addHandler(self.cap)
        self.plain.setLevel(logging.INFO)
        self.plain.propagate = False
        self._old = None

    def __enter__(self):
        self._old = es.logger
        es.logger = self.plain
        return self.cap

    def __exit__(self, *exc):
        es.logger = self._old
        return False


# ===========================================================================
# parse_accept_header (keys 10-13)
# ===========================================================================
def test_c35ak_01_wildcard_first_mime_decides_c35():
    # shipped returns on the FIRST token: the wildcard hits before json is read
    assert parse_accept_header("*/*, application/json") == ExportFormat.CSV
    assert parse_accept_header("text/*, application/json") == ExportFormat.CSV
    assert parse_accept_header("application/json, */*") == ExportFormat.JSON


def test_c35ak_02_accept_plain_shapes_c35():
    assert parse_accept_header(None) == ExportFormat.CSV
    assert parse_accept_header("") == ExportFormat.CSV
    assert parse_accept_header("text/csv;q=0.9") == ExportFormat.CSV
    assert parse_accept_header("application/vnd.ms-excel") == ExportFormat.EXCEL


# ===========================================================================
# get_selected_columns (key 6) - `or True` fallback kill
# ===========================================================================
def test_c35ak_03_all_invalid_falls_back_to_extended_c35():
    assert get_selected_columns(["nope"]) == EXTENDED_EXPORT_COLUMNS
    assert get_selected_columns(["nope", "nada"]) == EXTENDED_EXPORT_COLUMNS
    assert get_selected_columns(None) == EXTENDED_EXPORT_COLUMNS
    assert get_selected_columns([]) == EXTENDED_EXPORT_COLUMNS
    assert get_selected_columns(["event_id", "bogus"]) == [("event_id", "Event ID")]


# ===========================================================================
# filter_row_to_dict (9) / format_export_value (6)
# ===========================================================================
def test_c35ak_04_missing_field_reads_none_c35():
    row = EventExportRow(1, "cam", None, None, None, None, None, 0, False)
    del row.summary  # shipped 3-arg getattr -> None; the dropped-default RAISES
    out = filter_row_to_dict(row, ["event_id", "summary"])
    assert out == {"event_id": 1, "summary": None}
    assert format_export_value(row, "summary") == ""


def test_c35ak_05_format_export_value_shapes_c35():
    row = EventExportRow(1, "=bad", T0, None, 87, "high", "s", 3, True)
    assert format_export_value(row, "started_at") == "2026-01-05T10:30:45+00:00"
    assert format_export_value(row, "reviewed") == "Yes"
    row.reviewed = False
    assert format_export_value(row, "reviewed") == "No"
    assert format_export_value(row, "risk_score") == "87"
    assert format_export_value(row, "camera_name") == "'=bad"
    row.summary = None
    assert format_export_value(row, "summary") == ""


# ===========================================================================
# generate_export_filename (3)
# ===========================================================================
def test_c35ak_06_filename_timestamp_is_utc_c35():
    with TzPinned():
        name = generate_export_filename("events_export", ExportFormat.CSV)
        m = re.fullmatch(rf"events_export_({TS_RE})\.csv", name)
        assert m, name
        check_utc_timestamp(m.group(1))
        assert generate_export_filename("e", ExportFormat.EXCEL).endswith(".xlsx")


# ===========================================================================
# detections_to_json (10, 15, 17, 18)
# ===========================================================================
def test_c35ak_07_detections_json_bytes_and_missing_field_c35():
    det = DetectionExportRow(
        1,
        "cam",
        T0,
        "person",
        0.9,
        1,
        2,
        3,
        4,
        file_path="/a.jpg",
        thumbnail_path="/t.jpg",
        media_type="image",
    )
    expected = [
        {
            "detection_id": 1,
            "camera_name": "cam",
            "detected_at": T0.isoformat(),
            "object_type": "person",
            "confidence": 0.9,
            "bbox_x": 1,
            "bbox_y": 2,
            "bbox_width": 3,
            "bbox_height": 4,
            "file_path": "/a.jpg",
            "media_type": "image",
        }
    ]
    assert detections_to_json([det]) == json.dumps(expected, indent=2)
    # delete a field WITHOUT a dataclass default: the class-attribute default
    # would shadow the instance del and make the 2-arg getattr return None.
    del det.object_type  # shipped reads None; the dropped-default mutant RAISES
    expected[0]["object_type"] = None
    assert detections_to_json([det]) == json.dumps(expected, indent=2)


# ===========================================================================
# ExportService.__init__ mkdir (2, 4, 6)
# ===========================================================================
def test_c35ak_08_mkdir_creates_nested_c35():
    tmp = Path(tempfile.mkdtemp())
    target = tmp / "a" / "b"
    with patch.object(es, "EXPORT_DIR", target):
        ExportService(db=None)
    assert target.is_dir()  # parents True -> False/None: FileNotFoundError


def test_c35ak_09_mkdir_existing_dir_ok_c35():
    tmp = Path(tempfile.mkdtemp())
    target = tmp / "a" / "b"
    target.mkdir(parents=True)
    with patch.object(es, "EXPORT_DIR", target):
        ExportService(db=None)  # exist_ok dropped: FileExistsError


# ===========================================================================
# _create_empty_export (3, 18, 20, 24, 32, 34, 37, 50)
# ===========================================================================
def test_c35ak_10_empty_csv_c35():
    svc = ExportService(db=FakeSession([]))
    d = SpyDir()
    with TzPinned(), patch.object(es, "EXPORT_DIR", d):
        res = run(svc._create_empty_export("csv"))
    name = next(iter(d.files))
    check_utc_timestamp(filename_stamp(name, "csv"))
    content, kwargs = d.sole().last_text()
    assert kwargs == {"encoding": "utf-8"}
    assert content == ",".join(c[1] for c in EXTENDED_EXPORT_COLUMNS) + "\n"
    assert res == {
        "file_path": f"/api/exports/{name}",
        "file_size": len(content.encode("utf-8")),
        "event_count": 0,
        "format": "csv",
    }


def test_c35ak_11_empty_json_c35():
    svc = ExportService(db=FakeSession([]))
    d = SpyDir()
    with TzPinned(), patch.object(es, "EXPORT_DIR", d):
        res = run(svc._create_empty_export("json"))
    name = next(iter(d.files))
    check_utc_timestamp(filename_stamp(name, "json"))
    content, kwargs = d.sole().last_text()
    assert kwargs == {"encoding": "utf-8"}
    assert content == "[]"
    assert res["event_count"] == 0 and res["format"] == "json"
    assert res["file_path"] == f"/api/exports/{name}"


def test_c35ak_12_empty_zip_c35():
    svc = ExportService(db=FakeSession([]))
    d = SpyDir()
    fz = FakeZipModule()
    with TzPinned(), patch.object(es, "EXPORT_DIR", d), patch.object(es, "zipfile", fz):
        res = run(svc._create_empty_export("zip"))
    name = next(iter(d.files))
    stamp = filename_stamp(name, "zip")
    check_utc_timestamp(stamp)
    assert len(fz.records) == 1
    rec = fz.records[0]
    assert rec.ctor[0] == ("w", zipfile.ZIP_DEFLATED) and rec.ctor[1] == {}
    assert rec.entries == [(f"events_export_{stamp}.json", "[]")]
    assert res["format"] == "zip" and res["event_count"] == 0


def test_c35ak_13_empty_bad_format_c35():
    svc = ExportService(db=FakeSession([]))
    exc = assert_raises(ValueError, lambda: run(svc._create_empty_export("excel")))
    assert str(exc) == "Unsupported export format: excel"


# ===========================================================================
# export_events_with_progress
# ===========================================================================
def _progress_run(fmt, pins, tracker=None, columns=None, **filters):
    svc = ExportService(db=FakeSession(pins))
    tracker = tracker or Tracker()
    d = SpyDir()
    fz = FakeZipModule()
    with (
        LoggerSwap() as cap,
        TzPinned(),
        patch.object(es, "EXPORT_DIR", d),
        patch.object(es, "zipfile", fz),
    ):
        res = run(
            svc.export_events_with_progress("job-1", tracker, fmt, columns=columns, **filters)
        )
    return res, tracker, d, fz, cap


def test_c35ak_14_progress_db_none_exact_message_c35():
    svc = ExportService(db=None)
    exc = assert_raises(
        ValueError,
        lambda: run(svc.export_events_with_progress("j", Tracker(), "csv")),
    )
    assert str(exc) == "Database session required for export_events_with_progress"


def test_c35ak_15_progress_zero_events_exact_call_c35():
    pins = [(("count(*)",), 0, [])]
    res, tracker, d, _, _ = _progress_run("csv", pins)
    assert tracker.calls == [(("job-1", 50), {"message": "No events to export"})]
    name = next(iter(d.files))
    content, kwargs = d.sole().last_text()
    assert kwargs == {"encoding": "utf-8"}
    assert content == ",".join(c[1] for c in EXTENDED_EXPORT_COLUMNS) + "\n"
    assert res["event_count"] == 0 and res["format"] == "csv"


def _expected_csv(rows_text):
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow([c[1] for c in EXTENDED_EXPORT_COLUMNS])
    for r in rows_text:
        writer.writerow(r)
    return out.getvalue()


def test_c35ak_16_progress_csv_full_c35():
    pins = three_event_pins()
    res, tracker, d, _, cap = _progress_run("csv", pins)
    assert tracker.calls == [
        (("job-1", 10), {"message": "Found 3 events to export"}),
        (("job-1", 80), {"message": "Writing CSV file..."}),
        (("job-1", 95), {"message": "Finalizing export..."}),
    ]
    name = next(iter(d.files))
    check_utc_timestamp(filename_stamp(name, "csv"))
    content, kwargs = d.sole().last_text()
    assert kwargs == {"encoding": "utf-8"}
    expected = _expected_csv([BASE_ROW, unknown_row(2), unknown_row(3)])
    assert content == expected
    assert res == {
        "file_path": f"/api/exports/{name}",
        "file_size": len(expected.encode("utf-8")),
        "event_count": 3,
        "format": "csv",
    }
    assert len(cap.records) == 1
    rec = cap.records[0]
    assert rec.getMessage() == "Export completed"
    assert log_extras(rec) == {
        "job_id": "job-1",
        "format": "csv",
        "event_count": 3,
        "file_size": len(expected.encode("utf-8")),
    }


def test_c35ak_17_progress_csv_unknown_literals_c35():
    pins = three_event_pins()
    _, _, d, _, _ = _progress_run("csv", pins)
    content, _ = d.sole().last_text()
    lines = content.splitlines()
    assert lines[1].split(",")[1] == "Front Door"
    assert lines[2].split(",")[1] == "Unknown"
    assert lines[3].split(",")[1] == "Unknown"


def test_c35ak_18_progress_json_bytes_and_columns_c35():
    pins = three_event_pins()
    columns = ["event_id", "camera_name", "started_at", "summary"]
    _, _, d, _, _ = _progress_run("json", pins, columns=columns)
    name = next(iter(d.files))
    check_utc_timestamp(filename_stamp(name, "json"))
    content, kwargs = d.sole().last_text()
    assert kwargs == {"encoding": "utf-8"}
    expected = [
        {
            "event_id": eid,
            "camera_name": cname,
            "started_at": T0.isoformat(),
            "summary": "=SUM(1)",
        }
        for eid, cname in ((1, "Front Door"), (2, "Unknown"), (3, "Unknown"))
    ]
    assert content == json.dumps(expected, indent=2)


def test_c35ak_19_progress_zip_explicit_columns_c35():
    pins = three_event_pins()
    columns = ["event_id", "summary"]
    _, _, d, fz, _ = _progress_run("zip", pins, columns=columns)
    name = next(iter(d.files))
    stamp = filename_stamp(name, "zip")
    check_utc_timestamp(stamp)
    assert len(fz.records) == 1
    rec = fz.records[0]
    assert rec.ctor[0] == ("w", zipfile.ZIP_DEFLATED) and rec.ctor[1] == {}
    # the zip branch passes the CALLER's columns through filter_row_to_dict;
    # the `filter_row_to_dict(row, None)` mutant ignores them -> full rows
    expected = [
        {
            "event_id": eid,
            "summary": "=SUM(1)",
        }
        for eid in (1, 2, 3)
    ]
    assert rec.entries == [(f"events_export_{stamp}.json", json.dumps(expected, indent=2))]


def test_c35ak_20_progress_excel_bytes_c35():
    pins = three_event_pins()
    res, _, d, _, _ = _progress_run("excel", pins)
    name = next(iter(d.files))
    check_utc_timestamp(filename_stamp(name, "xlsx"))
    f = d.sole()
    assert len(f.buffers) == 1 and len(f.buffers[0]) > 1000
    assert res["file_size"] == len(f.buffers[0])


def test_c35ak_21_progress_tick_sequence_c35():
    rows = [ev(id=i) for i in range(1, 201)]
    pins = [
        (("count(*)",), 200, []),
        (("SELECT cameras.name", "cameras.id = 'cam-a'"), "Front Door", []),
        (("FROM events", "ORDER BY events.started_at DESC"), None, rows),
    ]
    tracker = Tracker()
    _progress_run("csv", pins, tracker=tracker)
    ticks = [c for c in tracker.calls if str(c[1].get("message", "")).startswith("Processing")]
    assert ticks == [
        (("job-1", 45), {"message": "Processing event 100/200"}),
        (("job-1", 80), {"message": "Processing event 200/200"}),
    ]
    assert tracker.calls[0] == (("job-1", 10), {"message": "Found 200 events to export"})


def test_c35ak_22_progress_startdate_mid_z_message_c35():
    pins = [(("count(*)",), 0, [])]
    exc = assert_raises(ValueError, lambda: _progress_run("csv", pins, start_date="2026Z-01-05"))
    assert "2026+00:00-01-05" in str(exc), f"message must carry replaced text: {exc}"


def test_c35ak_23_progress_startdate_lower_z_raises_c35():
    pins = [(("count(*)",), 0, [])]
    exc = assert_raises(
        ValueError,
        lambda: _progress_run("csv", pins, start_date="2026-01-05T00:00:00z"),
    )
    assert "'2026-01-05T00:00:00z'" in str(exc)


def test_c35ak_24_progress_enddate_z_variants_c35():
    pins = [(("count(*)",), 0, [])]
    exc = assert_raises(ValueError, lambda: _progress_run("csv", pins, end_date="2026Z-01-05"))
    assert "2026+00:00-01-05" in str(exc)
    exc = assert_raises(
        ValueError,
        lambda: _progress_run("csv", pins, end_date="2026-01-05T00:00:00z"),
    )
    assert "'2026-01-05T00:00:00z'" in str(exc)


def test_c35ak_25_progress_native_z_parses_c35():
    pins = [(("count(*)",), 0, [])]
    _, tracker, _, _, _ = _progress_run(
        "csv", pins, start_date="2026-01-05T00:00:00Z", end_date="2026-01-06T00:00:00Z"
    )
    assert tracker.calls[0][1] == {"message": "No events to export"}


# ===========================================================================
# export_events_with_websocket
# ===========================================================================
def _ws_run(fmt, pins, rep=None, **filters):
    rep = rep or Reporter()
    svc = ExportService(db=FakeSession(pins))
    d = SpyDir()
    fz = FakeZipModule()
    with (
        LoggerSwap() as cap,
        TzPinned(),
        patch.object(es, "EXPORT_DIR", d),
        patch.object(es, "zipfile", fz),
    ):
        res = run(svc.export_events_with_websocket(rep, fmt, **filters))
    return res, rep, d, fz, cap


def test_c35ak_26_ws_db_none_no_fail_call_c35():
    rep = Reporter()
    svc = ExportService(db=None)
    exc = assert_raises(ValueError, lambda: run(svc.export_events_with_websocket(rep, "csv")))
    assert str(exc) == "Database session required for export_events_with_websocket"
    assert rep.events == [] and rep.fails == []


def test_c35ak_27_ws_zero_path_c35():
    pins = [(("count(*)",), 0, [])]
    res, rep, d, _, _ = _ws_run("csv", pins)
    assert rep.events[0] == (
        "start",
        {
            "export_format": "csv",
            "filters": {
                "camera_id": None,
                "risk_level": None,
                "start_date": None,
                "end_date": None,
                "reviewed": None,
            },
        },
    )
    assert rep.events[1] == ("progress", 1, "Found 0 events to export", True)
    name = next(iter(d.files))
    content, kwargs = d.sole().last_text()
    assert kwargs == {"encoding": "utf-8"}
    assert rep.events[2][0] == "complete"
    assert rep.events[2][1] == {**res, "message": "No events to export"}
    assert res["event_count"] == 0


def test_c35ak_28_ws_csv_full_sequence_c35():
    pins = three_event_pins()
    res, rep, d, _, cap = _ws_run("csv", pins)
    kinds = [e[0] for e in rep.events]
    assert kinds == ["start"] + ["progress"] * 6 + ["complete"]
    assert rep.events[1] == ("progress", 1, "Found 3 events to export", True)
    assert rep.events[2] == ("progress", 23, "Processing event 1/3", False)
    assert rep.events[3] == ("progress", 46, "Processing event 2/3", False)
    assert rep.events[4] == ("progress", 70, "Processing event 3/3", False)
    assert rep.events[5] == ("progress", 80, "Writing CSV file...", True)
    assert rep.events[6] == ("progress", 95, "Finalizing export...", True)
    name = next(iter(d.files))
    check_utc_timestamp(filename_stamp(name, "csv"))
    content, kwargs = d.sole().last_text()
    assert kwargs == {"encoding": "utf-8"}
    assert content == _expected_csv([BASE_ROW, unknown_row(2), unknown_row(3)])
    assert res == {
        "file_path": f"/api/exports/{name}",
        "file_size": len(content.encode("utf-8")),
        "event_count": 3,
        "format": "csv",
        "duration_seconds": 1.25,
    }
    assert rep.events[7] == ("complete", res)
    assert len(cap.records) == 1
    rec = cap.records[0]
    assert rec.getMessage() == "Export completed with WebSocket events"
    assert log_extras(rec) == {
        "job_id": "job-ws-1",
        "format": "csv",
        "event_count": 3,
        "file_size": len(content.encode("utf-8")),
        "duration_seconds": 1.25,
    }


def _ws_inline(i, cname, started=T0.isoformat(), ended=T1.isoformat()):
    return {
        "event_id": i,
        "camera_name": cname,
        "started_at": started,
        "ended_at": ended,
        "risk_score": 87,
        "risk_level": "high",
        "summary": "=SUM(1)",
        "detection_count": 3,
        "reviewed": True,
        "object_types": "person,car",
        "reasoning": "because",
    }


def _ws_inline_expected():
    return [
        _ws_inline(1, "Front Door"),
        _ws_inline(2, "Unknown"),
        _ws_inline(3, "Unknown"),
    ]


def test_c35ak_29_ws_json_inline_dict_c35():
    pins = three_event_pins()
    _, _, d, _, _ = _ws_run("json", pins)
    content, kwargs = d.sole().last_text()
    assert kwargs == {"encoding": "utf-8"}
    assert content == json.dumps(_ws_inline_expected(), indent=2)


def test_c35ak_30_ws_zip_inline_dict_c35():
    pins = three_event_pins()
    _, _, d, fz, _ = _ws_run("zip", pins)
    assert len(fz.records) == 1
    rec = fz.records[0]
    assert rec.ctor[0] == ("w", zipfile.ZIP_DEFLATED) and rec.ctor[1] == {}
    assert rec.entries[0][1] == json.dumps(_ws_inline_expected(), indent=2)


def test_c35ak_31_ws_null_dates_both_writers_c35():
    rows = [ev(id=1, started_at=None, ended_at=None, object_types=None, reasoning=None)]
    pins = [
        (("count(*)",), 1, []),
        (("SELECT cameras.name", "cameras.id = 'cam-a'"), "Front Door", []),
        (("FROM events", "ORDER BY events.started_at DESC"), None, rows),
    ]
    _, _, d, _, _ = _ws_run("json", pins)
    got = json.loads(d.sole().last_text()[0])
    assert got[0]["started_at"] is None and got[0]["ended_at"] is None
    assert got[0]["object_types"] is None and got[0]["reasoning"] is None
    fz = FakeZipModule()
    rep = Reporter()
    svc = ExportService(db=FakeSession(pins))
    d2 = SpyDir()
    with patch.object(es, "EXPORT_DIR", d2), patch.object(es, "zipfile", fz):
        run(svc.export_events_with_websocket(rep, "zip"))
    zipped = json.loads(fz.records[0].entries[0][1])
    assert zipped[0]["started_at"] is None and zipped[0]["ended_at"] is None


def test_c35ak_32_ws_z_date_fail_path_c35():
    pins = [(("count(*)",), 0, [])]
    boom = ValueError("boom-z")
    rep = Reporter()
    svc = ExportService(db=FakeSession(pins))
    got = assert_raises(
        ValueError,
        lambda: run(svc.export_events_with_websocket(rep, "csv", start_date="2026Z-01-05")),
    )
    assert got is not boom
    assert "2026+00:00-01-05" in str(got)
    assert len(rep.fails) == 1
    assert rep.fails[0][0] is got
    assert rep.fails[0][1] is False
    exc = assert_raises(ValueError, lambda: _ws_run("csv", pins, end_date="2026Z-01-05"))
    assert "2026+00:00-01-05" in str(exc), f"end_date replaced text missing: {exc}"


def test_c35ak_33_ws_lower_z_raises_c35():
    pins = [(("count(*)",), 0, [])]
    exc = assert_raises(
        ValueError,
        lambda: _ws_run("csv", pins, end_date="2026-01-05T00:00:00z"),
    )
    assert "'2026-01-05T00:00:00z'" in str(exc)


def test_c35ak_34_ws_start_failure_fails_job_c35():
    pins = [(("count(*)",), 0, [])]
    boom = RuntimeError("boom")
    rep = Reporter(start_raises=boom)
    got = assert_raises(
        RuntimeError,
        lambda: run(ExportService(db=FakeSession(pins)).export_events_with_websocket(rep, "csv")),
    )
    assert got is boom
    assert rep.fails == [(boom, False)]


# ===========================================================================
# events_to_excel
# ===========================================================================
def _excel(data):
    return openpyxl.load_workbook(io.BytesIO(data)).active


def _one_row_sheet():
    row = EventExportRow(
        1,
        "Front-Door-Camera-Long",
        T0,
        None,
        87,
        "high",
        "=SUM(1)",
        3,
        True,
        "person,car",
        "because",
    )
    return events_to_excel([row])


def test_c35ak_35_excel_header_and_body_styles_c35():
    ws = _excel(_one_row_sheet())
    cell = ws["A1"]
    assert cell.value == "Event ID"
    assert cell.font.bold is True
    assert cell.font.color is not None and cell.font.color.rgb == "00FFFFFF"
    assert cell.fill.patternType == "solid"
    assert cell.fill.start_color.rgb == "004472C4"
    assert cell.fill.end_color.rgb == "004472C4"
    assert cell.alignment.horizontal == "center"
    assert cell.alignment.vertical == "center"
    assert cell.alignment.wrap_text is True
    for side in ("left", "right", "top", "bottom"):
        b = getattr(cell.border, side)
        assert b is not None and b.style == "thin", f"header {side}"
    body = ws["A2"]
    for side in ("left", "right", "top", "bottom"):
        b = getattr(body.border, side)
        assert b is not None and b.style == "thin", f"body {side}"


def test_c35ak_36_excel_alt_row_fills_c35():
    rows = [EventExportRow(i, f"c{i}", T0, None, None, None, None, 0, False) for i in (1, 2, 3)]
    ws = _excel(events_to_excel(rows))
    assert ws["A2"].fill.patternType == "solid"
    assert ws["A2"].fill.start_color.rgb == "00E9EDF5"
    assert ws["A2"].fill.end_color.rgb == "00E9EDF5"
    assert ws["A3"].fill.start_color.rgb != "00E9EDF5"
    assert ws["A4"].fill.start_color.rgb == "00E9EDF5"


def test_c35ak_37_excel_widths_follow_content_c35():
    ws = _excel(_one_row_sheet())
    dims = ws.column_dimensions
    assert dims["A"].width == 10.0  # header seed len("Event ID") + 2
    assert dims["B"].width == 24.0  # "Front-Door-Camera-Long" + 2
    # (kills value=None and format_export_value(None, f): widths stay header-
    # seeded, so column B sits at 8.0)


def test_c35ak_38_excel_cap_and_seed_zero_c35():
    long_row = EventExportRow("X" * 60, "c", None, None, None, None, "S" * 60, 0, False)
    ws = _excel(events_to_excel([long_row], columns=[("event_id", ""), ("summary", "Summary")]))
    assert ws.column_dimensions["A"].width == 52.0  # min(60, 50) + 2
    assert ws.column_dimensions["B"].width == 52.0
    zero = EventExportRow(None, "c", None, None, None, None, None, 0, False)
    ws2 = _excel(events_to_excel([zero], columns=[("event_id", "")]))
    # seed 0, falsy body: shipped keeps 0 + 2; the else-1 mutant raises it to 3
    assert ws2.column_dimensions["A"].width == 2.0


def test_c35ak_39_excel_tie_above_cap_c35():
    row = EventExportRow("c", "Z" * 55, None, None, None, None, None, 0, False)
    ws = _excel(events_to_excel([row], columns=[("camera_name", "Y" * 55)]))
    # header-seed 55 dominates; shipped KEEPS the 55 (tie is not >), the >=
    # mutant stores min(55, 50) = 50 on the tie
    assert ws.column_dimensions["A"].width == 57.0


def test_c35ak_40_excel_freeze_panes_c35():
    ws = _excel(_one_row_sheet())
    assert ws.freeze_panes == "A2"


def test_c35ak_41_excel_none_cell_is_empty_string_c35():
    # the `raw_value is None` branch writes ""; the m124 mutant writes None.
    # event_id=None takes that branch ("" would take the sanitize branch).
    # openpyxl VALUE readback cannot tell them apart - the sheet XML can:
    # "" -> <c t="inlineStr"></c>, None -> value-less <c/>.
    zero = EventExportRow(None, "c", None, None, None, None, None, 0, False)
    data = events_to_excel([zero], columns=[("event_id", "Event ID")])
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        xml = z.read("xl/worksheets/sheet1.xml").decode()
    body = re.search(r"<row [^>]*r=\"2\".*?</row>", xml, re.S).group(0)
    assert 't="inlineStr"' in body, f"None-branch cells must be empty strings: {body}"


def test_c35ak_42_excel_missing_field_reads_none_c35():
    row = EventExportRow(1, "c", None, None, None, None, "s", 0, False)
    del row.summary  # NO class default: shipped 3-arg getattr -> None; 2-arg RAISES
    ws = _excel(events_to_excel([row], columns=[("summary", "Summary")]))
    assert ws["A2"].value is None  # shipped renders "" -> reloads as None inlineStr


def test_c35ak_44_excel_guard_uses_own_column_c35():
    # m148 compares the cell length against the NEIGHBOUR's width. Unequal
    # seeds 6 vs 2 and a 4-char body cell only pass their OWN guard: shipped
    # raises B to 4+2=6, the neighbour-compare mutant leaves it at 2+2=4.
    row = EventExportRow("X", "c", None, None, 1234, None, None, 0, False)
    ws = _excel(events_to_excel([row], columns=[("event_id", "123456"), ("risk_score", "12")]))
    assert ws.column_dimensions["A"].width == 8.0
    assert ws.column_dimensions["B"].width == 6.0  # mutant: 4.0


def test_c35ak_43_excel_bool_cells_c35():
    row = EventExportRow(1, "c", None, None, None, None, None, 0, True)
    ws = _excel(events_to_excel([row]))
    assert ws["I2"].value == "Yes"
    row2 = EventExportRow(1, "c", None, None, None, None, None, 0, False)
    ws2 = _excel(events_to_excel([row2]))
    assert ws2["I2"].value == "No"
